"""Build the Word test matrix from the reviewable Markdown table."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "UNIT_TEST_TABLE.md"
OUTPUT = ROOT / "docs" / "UNIT_TEST_TABLE.docx"
HEADERS = [
    "Proponent",
    "Module Name",
    "Unit Name",
    "Date Tested",
    "Test Case ID",
    "Test Case Description",
    "Expected Result",
    "Actual Result",
    "Result",
]
WIDTHS = [0.76, 0.94, 0.93, 0.70, 0.80, 1.85, 2.05, 1.02, 0.75]


def read_rows() -> list[list[str]]:
    lines = [line for line in SOURCE.read_text(encoding="utf-8").splitlines() if line.startswith("|")]
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in lines]
    if rows[0] != HEADERS or len(rows[0]) != len(WIDTHS):
        raise ValueError("The table headings do not match the report template")
    data = rows[2:]
    if len(data) != 63:
        raise ValueError(f"Expected 63 cases, found {len(data)}")
    for number, row in enumerate(data, 1):
        if len(row) != len(HEADERS) or row[4] != f"CASE-{number:03d}":
            raise ValueError(f"Invalid case row {number}")
        if row[0] or row[3]:
            raise ValueError("Proponent and Date Tested must remain blank")
    return data


def set_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", "80"), ("bottom", "80"), ("left", "65"), ("right", "65")):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tc_pr.append(margins)


def set_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tc_pr.append(borders)


def set_row_property(row, property_name: str) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement(f"w:{property_name}"))


def fill_cell(cell, text: str, column: int, header: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_margins(cell)
    set_borders(cell)
    if header:
        set_shading(cell, "D9EAD3")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if header or column in (0, 3, 4, 8) else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8.5 if header else 8)
        run.font.bold = header
        run.font.color.rgb = RGBColor(0, 0, 0)


def build() -> None:
    data = read_rows()
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    title_border = OxmlElement("w:pBdr")
    bottom_border = OxmlElement("w:bottom")
    bottom_border.set(qn("w:val"), "nil")
    title_border.append(bottom_border)
    title._p.get_or_add_pPr().append(title_border)
    run = title.add_run("Unit Testing")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(8)
    run = note.add_run("Jest and React Native Testing Library  |  63 automated cases")
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0, 0, 0)

    table = document.add_table(rows=1, cols=len(HEADERS))
    table.autofit = False
    for column, width in zip(table.columns, WIDTHS):
        column.width = Inches(width)
    for index, heading in enumerate(HEADERS):
        table.rows[0].cells[index].width = Inches(WIDTHS[index])
        fill_cell(table.rows[0].cells[index], heading, index, header=True)
    set_row_property(table.rows[0], "tblHeader")
    set_row_property(table.rows[0], "cantSplit")

    for source_row in data:
        row = table.add_row()
        set_row_property(row, "cantSplit")
        for index, value in enumerate(source_row):
            row.cells[index].width = Inches(WIDTHS[index])
            fill_cell(row.cells[index], value, index)

    document.save(OUTPUT)
    print(f"Created {OUTPUT} with {len(data)} cases")


if __name__ == "__main__":
    build()
